from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _parse_hex_color(value: str) -> str | None:
    """
    Returns ARGB hex (AARRGGBB) for openpyxl or None.
    Accepts #RRGGBB, #RGB, rgb(r,g,b).
    """
    if not value:
        return None
    v = value.strip().lower()
    if v.startswith("#"):
        h = v[1:]
        if len(h) == 3:
            h = "".join([c * 2 for c in h])
        if len(h) == 6 and re.fullmatch(r"[0-9a-f]{6}", h):
            return "FF" + h.upper()
        return None
    m = re.match(r"rgb\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*\)", v)
    if m:
        r, g, b = [max(0, min(255, int(x))) for x in m.groups()]
        return f"FF{r:02X}{g:02X}{b:02X}"
    return None


def _extract_workbook(content: Dict[str, Any]) -> Tuple[List[Dict[str, Any]], int]:
    """
    content: model.content (Document/SectionTable) dict
    returns (sheets, active_index)
    """
    if not isinstance(content, dict):
        return [], 0
    cs = content.get("custom_sheet")
    if not isinstance(cs, dict):
        return [], 0
    if isinstance(cs.get("sheets"), list) and cs["sheets"]:
        ai = cs.get("activeSheetIndex") or 0
        try:
            ai = int(ai)
        except Exception:
            ai = 0
        ai = max(0, min(ai, len(cs["sheets"]) - 1))
        return cs["sheets"], ai
    if isinstance(cs.get("data"), list):
        return [cs], 0
    return [], 0


def export_custom_sheet_to_xlsx_bytes(
    title: str,
    content: Dict[str, Any],
) -> bytes:
    sheets, _ = _extract_workbook(content or {})
    wb = Workbook()

    # Remove default sheet; we'll recreate
    if wb.worksheets:
        wb.remove(wb.worksheets[0])

    if not sheets:
        sheets = [{"name": "Лист1", "data": [[""]]}]

    for idx, sh in enumerate(sheets):
        name = str(sh.get("name") or f"Лист{idx + 1}")[:31] or f"Лист{idx + 1}"
        ws = wb.create_sheet(title=name)
        data = sh.get("data") if isinstance(sh.get("data"), list) else []
        styles = sh.get("styles") if isinstance(sh.get("styles"), dict) else {}
        col_widths = (
            sh.get("colWidths") if isinstance(sh.get("colWidths"), list) else []
        )
        row_heights = (
            sh.get("rowHeights") if isinstance(sh.get("rowHeights"), list) else []
        )

        # Column widths (px -> approx excel width)
        for c, px in enumerate(col_widths, start=1):
            try:
                w = float(px)
            except Exception:
                continue
            ws.column_dimensions[get_column_letter(c)].width = max(
                6.0, min(60.0, w / 7.0)
            )

        # Row heights (px -> points)
        for r, px in enumerate(row_heights, start=1):
            try:
                h = float(px)
            except Exception:
                continue
            ws.row_dimensions[r].height = max(12.0, min(120.0, h * 0.75))

        for r_idx, row in enumerate(data, start=1):
            if not isinstance(row, list):
                continue
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(
                    row=r_idx, column=c_idx, value="" if val is None else str(val)
                )
                sk = f"{r_idx - 1}:{c_idx - 1}"
                st = styles.get(sk) if isinstance(styles, dict) else None
                if not st:
                    continue
                # Font
                font_kwargs = {}
                if st.get("bold") is True:
                    font_kwargs["bold"] = True
                if st.get("italic") is True:
                    font_kwargs["italic"] = True
                ff = st.get("fontFamily")
                if isinstance(ff, str) and ff.strip():
                    font_kwargs["name"] = ff.strip().strip('"').strip("'")
                fs = st.get("fontSize")
                try:
                    fs_n = float(fs)
                    if 6 <= fs_n <= 72:
                        font_kwargs["size"] = fs_n
                except Exception:
                    pass
                tc = _parse_hex_color(str(st.get("textColor") or ""))
                if tc:
                    font_kwargs["color"] = tc
                if font_kwargs:
                    cell.font = Font(**font_kwargs)

                # Alignment
                align = st.get("align")
                if align in ("left", "center", "right"):
                    cell.alignment = Alignment(horizontal=align)

                # Fill
                fc = _parse_hex_color(str(st.get("fillColor") or ""))
                if fc:
                    cell.fill = PatternFill("solid", fgColor=fc)

    out = io.BytesIO()
    wb.properties.title = title or "CNC Office"
    wb.save(out)
    return out.getvalue()
