from io import BytesIO

import pytest


def test_generate_monitoring_report_docx_contains_direction_and_remark(settings):
    pytest.importorskip("docx")
    from docx import Document

    from bot_api.reports import generate_monitoring_report_docx

    settings.BOT_SHEET_FIO_COL = 2
    settings.BOT_SHEET_GROUP_COL = 1
    settings.BOT_SHEET_REMARK_COL = 11
    content = {
        "custom_sheet": {
            "version": 2,
            "activeSheetIndex": 0,
            "sheets": [
                {
                    "name": "BIM",
                    "data": [
                        [
                            "",
                            "Группа",
                            "ФИО",
                            "Номер телефона",
                            "Телеграм",
                            "Вконтакте",
                            "ТикТок",
                            "Ответственный",
                            "01.04-05.04",
                            "06.04-12.04",
                            "13.04-19.04",
                            "20.04-26.04",
                            "Состояние",
                        ],
                        [
                            "",
                            "BIM-25-1",
                            "Иванов Иван",
                            "",
                            "@student",
                            "",
                            "",
                            "",
                            "",
                            "",
                            "",
                            "фото с оружием",
                            "учится",
                        ],
                    ],
                }
            ],
        }
    }

    blob = generate_monitoring_report_docx(
        table_title="Rangers",
        content=content,
        sheet_name="BIM",
        report_year=2026,
        report_month=4,
    )

    doc = Document(BytesIO(blob))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "BIM" in text
    assert "Всего проверено: 1 студент" in text
    assert "Иванов Иван" in text
    assert "фото с оружием" in text
    assert "Отчёт за апрель 2026 г." in text


def test_generate_monitoring_report_docx_uses_three_week_columns(settings):
    pytest.importorskip("docx")
    from docx import Document

    from bot_api.reports import generate_monitoring_report_docx

    settings.BOT_SHEET_FIO_COL = 2
    settings.BOT_SHEET_GROUP_COL = 1
    settings.BOT_SHEET_REMARK_COL = 11
    content = {
        "custom_sheet": {
            "version": 2,
            "activeSheetIndex": 0,
            "sheets": [
                {
                    "name": "ЧПУ",
                    "data": [
                        [
                            "",
                            "Группа",
                            "ФИО",
                            "Телеграм",
                            "04.05-10.05",
                            "11.05-17.05",
                            "18.05-24.05",
                            "Состояние",
                        ],
                        [
                            "",
                            "ЧПУ 1",
                            "Петров Пётр",
                            "@petrov",
                            "",
                            "опоздание",
                            "",
                            "учится",
                        ],
                    ],
                }
            ],
        }
    }

    blob = generate_monitoring_report_docx(
        table_title="Rangers",
        content=content,
        sheet_name="ЧПУ",
        report_year=2026,
        report_month=5,
    )

    doc = Document(BytesIO(blob))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "опоздание" in text
    assert "Петров Пётр" in text
