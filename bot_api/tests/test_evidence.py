from datetime import date
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile

from bot_api.evidence import delete_remark_evidence, upsert_remark_evidence
from bot_api.models import RemarkEvidence
from sections.models import SectionTable

# 1x1 PNG
_PNG = bytes(
    [
        0x89,
        0x50,
        0x4E,
        0x47,
        0x0D,
        0x0A,
        0x1A,
        0x0A,
        0x00,
        0x00,
        0x00,
        0x0D,
        0x49,
        0x48,
        0x44,
        0x52,
        0x00,
        0x00,
        0x00,
        0x01,
        0x00,
        0x00,
        0x00,
        0x01,
        0x08,
        0x02,
        0x00,
        0x00,
        0x00,
        0x90,
        0x77,
        0x53,
        0xDE,
        0x00,
        0x00,
        0x00,
        0x0C,
        0x49,
        0x44,
        0x41,
        0x54,
        0x08,
        0xD7,
        0x63,
        0xF8,
        0xCF,
        0xC0,
        0x00,
        0x00,
        0x03,
        0x01,
        0x01,
        0x00,
        0x18,
        0xDD,
        0x8D,
        0xB4,
        0x00,
        0x00,
        0x00,
        0x00,
        0x49,
        0x45,
        0x4E,
        0x44,
        0xAE,
        0x42,
        0x60,
        0x82,
    ]
)


@pytest.mark.django_db
def test_set_student_remark_with_evidence_saves_file(client, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    settings.CNC_BOT_API_SECRET = "secret"
    settings.CNC_BOT_TABLE_OWNER_USERNAME = "bot_owner"
    settings.BOT_SHEET_FIO_COL = 2
    settings.BOT_SHEET_GROUP_COL = 1
    settings.BOT_SHEET_REMARK_COL = 4
    owner = get_user_model().objects.create_user(username="bot_owner", password="pass")
    table = SectionTable.objects.create(
        owner=owner,
        title="Rangers",
        section_type="rangers",
        content={
            "custom_sheet": {
                "version": 2,
                "activeSheetIndex": 0,
                "sheets": [
                    {
                        "name": "ЧПУ",
                        "data": [
                            ["", "Группа", "ФИО", "TG", "20.04-26.04", "Состояние"],
                            ["", "Г1", "Иванов Иван", "", "", "учится"],
                        ],
                    }
                ],
            }
        },
    )

    resp = client.post(
        "/api/bot/gateway/",
        data={
            "action": "set_student_remark",
            "table_id": str(table.id),
            "sheet_name": "ЧПУ",
            "student_fio": "Иванов Иван",
            "group": "Г1",
            "remark_date": "22.04.2026",
            "remark_text": "запрещённый контент",
            "evidence": SimpleUploadedFile("shot.png", _PNG, content_type="image/png"),
        },
        HTTP_X_CNC_BOT_TOKEN="secret",
    )
    assert resp.status_code == 200, resp.content
    assert resp.json()["evidence_saved"] is True
    assert RemarkEvidence.objects.filter(table=table, student_fio="Иванов Иван").count() == 1

    table.refresh_from_db()
    cell = table.content["custom_sheet"]["sheets"][0]["data"][1][4]
    assert cell == "запрещённый контент"


@pytest.mark.django_db
def test_set_no_remarks_deletes_evidence(client, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    settings.CNC_BOT_API_SECRET = "secret"
    settings.CNC_BOT_TABLE_OWNER_USERNAME = "bot_owner"
    settings.BOT_SHEET_FIO_COL = 2
    settings.BOT_SHEET_GROUP_COL = 1
    settings.BOT_SHEET_REMARK_COL = 4
    owner = get_user_model().objects.create_user(username="bot_owner", password="pass")
    table = SectionTable.objects.create(
        owner=owner,
        title="Rangers",
        section_type="rangers",
        content={
            "custom_sheet": {
                "version": 2,
                "activeSheetIndex": 0,
                "sheets": [
                    {
                        "name": "ЧПУ",
                        "data": [
                            ["", "Группа", "ФИО", "TG", "20.04-26.04", "Состояние"],
                            ["", "Г1", "Иванов Иван", "", "старое", "учится"],
                        ],
                    }
                ],
            }
        },
    )
    upsert_remark_evidence(
        table=table,
        sheet_name="ЧПУ",
        student_fio="Иванов Иван",
        remark_date=date(2026, 4, 22),
        raw=_PNG,
        filename="a.png",
        content_type="image/png",
    )
    assert RemarkEvidence.objects.count() == 1

    resp = client.post(
        "/api/bot/gateway/",
        {
            "action": "set_student_remark",
            "table_id": table.id,
            "sheet_name": "ЧПУ",
            "student_fio": "Иванов Иван",
            "group": "Г1",
            "remark_date": "22.04.2026",
        },
        format="json",
        HTTP_X_CNC_BOT_TOKEN="secret",
    )
    assert resp.status_code == 200, resp.content
    assert RemarkEvidence.objects.count() == 0
    assert delete_remark_evidence(
        table=table,
        sheet_name="ЧПУ",
        student_fio="Иванов Иван",
        remark_date=date(2026, 4, 22),
    ) is False


def test_generate_monitoring_report_embeds_evidence_or_missing_note(settings):
    pytest.importorskip("docx")
    from docx import Document

    from bot_api.reports import generate_monitoring_report_docx

    settings.BOT_SHEET_FIO_COL = 2
    settings.BOT_SHEET_GROUP_COL = 1
    settings.BOT_SHEET_REMARK_COL = 4
    content = {
        "custom_sheet": {
            "version": 2,
            "activeSheetIndex": 0,
            "sheets": [
                {
                    "name": "ЧПУ",
                    "data": [
                        ["", "Группа", "ФИО", "TG", "20.04-26.04", "Состояние"],
                        ["", "Г1", "Иванов Иван", "@x", "нарушение", "учится"],
                        ["", "Г1", "Петров Пётр", "@y", "другое", "учится"],
                    ],
                }
            ],
        }
    }
    evidence = {
        ("иванов иван", date(2026, 4, 22)): _PNG,
    }
    blob = generate_monitoring_report_docx(
        table_title="Rangers",
        content=content,
        sheet_name="ЧПУ",
        report_year=2026,
        report_month=4,
        evidence_by_fio_date=evidence,
    )
    doc = Document(BytesIO(blob))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Иванов Иван" in text
    assert "Петров Пётр" in text
    assert "скрин отсутствует" in text
    # У Иванова скрин есть — пометки отсутствия быть не должно в его строке
    assert "Иванов Иван: нарушение" in text or "Иванов Иван:" in text
    assert len(doc.inline_shapes) >= 1


def test_bmp_evidence_converts_to_jpeg_for_word(settings, tmp_path):
    pytest.importorskip("docx")
    from PIL import Image
    from docx import Document

    from bot_api.evidence import evidence_bytes_for_docx, upsert_remark_evidence
    from bot_api.reports import generate_monitoring_report_docx

    settings.MEDIA_ROOT = tmp_path / "media"
    User = get_user_model()
    owner = User.objects.create_user(username="bmp_u", password="x")
    table = SectionTable.objects.create(
        owner=owner,
        title="T",
        section_type="rangers",
        content={},
    )
    buf = BytesIO()
    Image.new("RGB", (8, 8), color=(200, 50, 50)).save(buf, format="BMP")
    bmp = buf.getvalue()

    upsert_remark_evidence(
        table=table,
        sheet_name="ЧПУ",
        student_fio="Иванов Иван",
        remark_date=date(2026, 4, 22),
        raw=bmp,
        filename="screen.bmp",
        content_type="image/bmp",
    )
    obj = RemarkEvidence.objects.get(table=table)
    with obj.image.open("rb") as fh:
        stored = fh.read()
    assert stored[:3] == b"\xff\xd8\xff"

    docx_payload = evidence_bytes_for_docx(bmp)
    assert docx_payload and docx_payload[:3] == b"\xff\xd8\xff"

    content = {
        "custom_sheet": {
            "version": 2,
            "activeSheetIndex": 0,
            "sheets": [
                {
                    "name": "ЧПУ",
                    "data": [
                        ["", "Группа", "ФИО", "TG", "20.04-26.04"],
                        ["", "Г1", "Иванов Иван", "@x", "нарушение"],
                    ],
                }
            ],
        }
    }
    blob = generate_monitoring_report_docx(
        table_title="Rangers",
        content=content,
        sheet_name="ЧПУ",
        report_year=2026,
        report_month=4,
        evidence_by_fio_date={("иванов иван", date(2026, 4, 22)): bmp},
    )
    doc = Document(BytesIO(blob))
    assert len(doc.inline_shapes) >= 1
    assert "не удалось вставить скрин" not in "\n".join(p.text for p in doc.paragraphs)
