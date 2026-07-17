from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any, Optional

from django.conf import settings

from bot_api.monthly_rollover import detect_sheet_month, find_week_columns
from bot_api.sheet_utils import find_sheet_by_name, get_workbook_sheets
from bot_api.week_column import pick_week_col_by_date

_MONTH_NAMES_RU = (
    "",
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
)

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
except ImportError:  # pragma: no cover
    Document = None  # type: ignore
    WD_ALIGN_PARAGRAPH = None  # type: ignore
    Cm = None  # type: ignore
    Pt = None  # type: ignore


def generate_monitoring_report_docx(
    *,
    table_title: str,
    content: dict,
    sheet_name: str,
    report_date: Optional[date] = None,
    report_year: Optional[int] = None,
    report_month: Optional[int] = None,
    evidence_by_fio_date: Optional[dict[tuple[str, date], bytes]] = None,
) -> bytes:
    if Document is None:
        raise RuntimeError("python-docx не установлен")

    sheets, _ = get_workbook_sheets(content if isinstance(content, dict) else {})
    sheet = find_sheet_by_name(sheets, sheet_name)
    if not sheet:
        raise ValueError("Лист не найден")

    data = sheet.get("data") if isinstance(sheet.get("data"), list) else []
    header = data[0] if data and isinstance(data[0], list) else []
    direction = str(sheet.get("name") or sheet_name or table_title or "направление")
    fio_col = int(getattr(settings, "BOT_SHEET_FIO_COL", 2))
    group_col = int(getattr(settings, "BOT_SHEET_GROUP_COL", 1))
    fallback_remark_col = int(getattr(settings, "BOT_SHEET_REMARK_COL", 11))
    remark_cols = find_week_columns(data, fallback_remark_col)

    rows = _collect_report_rows(
        data,
        header,
        fio_col,
        group_col,
        remark_cols,
        evidence_by_fio_date or {},
    )
    groups = sorted({row["group"] for row in rows["students"] if row["group"]}, key=str.lower)
    findings = rows["findings"]

    period_year, period_month = _resolve_report_period(
        data,
        report_date=report_date,
        report_year=report_year,
        report_month=report_month,
    )

    doc = Document()
    _apply_default_font(doc)
    _build_cover(doc, direction, period_year, period_month)
    doc.add_page_break()
    _build_body(
        doc,
        direction,
        groups,
        rows["checked_count"],
        findings,
        period_year,
        period_month,
    )
    _normalize_doc_fonts(doc)

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def report_filename(sheet_name: str, year: Optional[int] = None, month: Optional[int] = None) -> str:
    period = f"_{month:02d}_{year}" if year and month else ""
    safe_sheet = _safe_filename(sheet_name or "report")
    return f"monitoring_{safe_sheet}{period}.docx"


def _resolve_report_period(
    sheet_data: list,
    *,
    report_date: Optional[date],
    report_year: Optional[int],
    report_month: Optional[int],
) -> tuple[int, int]:
    if report_year and report_month:
        return int(report_year), int(report_month)
    if report_date:
        return report_date.year, report_date.month
    detected = detect_sheet_month(sheet_data)
    if detected:
        return detected
    today = date.today()
    return today.year, today.month


def _format_report_period(year: int, month: int) -> str:
    if 1 <= month <= 12:
        return f"Отчёт за {_MONTH_NAMES_RU[month]} {year} г."
    return f"{year} г."


def _build_cover(doc, direction: str, report_year: int, report_month: int) -> None:
    for _ in range(6):
        doc.add_paragraph("")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "Проведен мониторинг социальных сетей студентов,\n"
        f"направления «{direction}», с целью выявления\n"
        "деструктивных элементов"
    )
    run.bold = True
    run.font.size = Pt(20)

    for _ in range(7):
        doc.add_paragraph("")

    developer = getattr(settings, "MONITORING_REPORT_DEVELOPER", "") or ""
    approver = getattr(settings, "MONITORING_REPORT_APPROVER", "") or ""

    p = doc.add_paragraph()
    p.add_run("Разработчик:\nПреподаватель (специалист)")
    p.add_run("\t\t\t\t\t\t")
    p.add_run(developer or "________________________")

    p = doc.add_paragraph()
    p.add_run("Согласовано:\nРуководитель проекта 2 уровня")
    p.add_run("\t\t\t\t\t")
    p.add_run(approver or "________________________")

    doc.add_paragraph("")
    period = doc.add_paragraph(_format_report_period(report_year, report_month))
    period.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _build_body(
    doc,
    direction: str,
    groups: list[str],
    checked_count: int,
    findings: list[dict],
    report_year: int,
    report_month: int,
) -> None:
    groups_text = ", ".join(groups) if groups else "группы не указаны"
    doc.add_paragraph(_format_report_period(report_year, report_month))
    doc.add_paragraph(
        f"Проведён просмотр открытых профилей и публикаций студентов групп: {groups_text}."
    )
    doc.add_paragraph(f"Всего проверено: {checked_count} {_student_word(checked_count)}.")
    doc.add_paragraph("Студенты, отнесённые к группе риска (по результатам мониторинга):")

    if not findings:
        doc.add_paragraph("Студенты с замечаниями не выявлены.")
        return

    current_group = None
    index = 1
    for item in findings:
        if item["group"] != current_group:
            current_group = item["group"]
            heading = doc.add_paragraph()
            run = heading.add_run(current_group or "Группа не указана")
            run.bold = True
            run.font.size = Pt(14)
            index = 1

        remark = item["remark"]
        social = f" - {item['social']}" if item.get("social") else ""
        missing = "" if item.get("evidence_bytes") else " (скрин отсутствует)"
        doc.add_paragraph(
            f"{index}. {item['fio']}: {remark}{social}{missing}",
            style=None,
        )
        if item.get("evidence_bytes"):
            _add_evidence_picture(doc, item["evidence_bytes"])
        index += 1


def _add_evidence_picture(doc, raw: bytes) -> None:
    try:
        stream = BytesIO(raw)
        doc.add_picture(stream, width=Cm(14))
    except Exception:
        note = doc.add_paragraph()
        note.add_run("(не удалось вставить скрин в документ)")


def _collect_report_rows(
    data: list,
    header: list,
    fio_col: int,
    group_col: int,
    remark_cols: list[int],
    evidence_by_fio_date: dict[tuple[str, date], bytes],
) -> dict[str, Any]:
    students = []
    findings = []
    for row in data[1:]:
        if not isinstance(row, list):
            continue
        fio = _cell(row, fio_col)
        if not fio:
            continue
        group = _cell(row, group_col)
        students.append({"fio": fio, "group": group})
        social = _first_social_link(row, header)
        for col in remark_cols:
            value = _cell(row, col)
            if not value or _is_no_remark(value):
                continue
            label = _cell(header, col)
            remark_text = f"{label}: {value}" if label else value
            evidence_bytes = _match_evidence_for_week(
                data,
                fio,
                col,
                evidence_by_fio_date,
            )
            findings.append(
                {
                    "fio": fio,
                    "group": group,
                    "remark": remark_text,
                    "social": social,
                    "week_col": col,
                    "evidence_bytes": evidence_bytes,
                }
            )
    findings.sort(key=lambda item: (item["group"].lower(), item["fio"].lower(), item["week_col"]))
    return {"students": students, "checked_count": len(students), "findings": findings}


def _match_evidence_for_week(
    sheet_data: list,
    fio: str,
    week_col: int,
    evidence_by_fio_date: dict[tuple[str, date], bytes],
) -> Optional[bytes]:
    fio_key = " ".join(fio.split()).strip().lower()
    for (ev_fio, ev_date), blob in evidence_by_fio_date.items():
        if ev_fio != fio_key:
            continue
        matched_col = pick_week_col_by_date(sheet_data, ev_date)
        if matched_col == week_col:
            return blob
    return None


def _first_social_link(row: list, header: list) -> str:
    for idx, name in enumerate(header):
        lowered = str(name or "").strip().lower()
        if not any(token in lowered for token in ("телеграм", "telegram", "тг", "вконтакте", "vk", "тик", "tiktok")):
            continue
        value = _cell(row, idx)
        if value:
            return value
    return ""


def _cell(row: list, idx: int) -> str:
    if idx < 0 or idx >= len(row):
        return ""
    return "" if row[idx] is None else str(row[idx]).strip()


def _is_no_remark(value: str) -> bool:
    lowered = value.strip().lower()
    if lowered in {"замечаний нет", "нет", "без замечаний", "-", "—"}:
        return True
    return lowered.startswith("замечаний нет")


def _student_word(count: int) -> str:
    if count % 10 == 1 and count % 100 != 11:
        return "студент"
    if count % 10 in (2, 3, 4) and count % 100 not in (12, 13, 14):
        return "студента"
    return "студентов"


def _safe_filename(value: str) -> str:
    return "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in value)[:80]


def _apply_default_font(doc) -> None:
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)


def _normalize_doc_fonts(doc) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            if run.font.size is None:
                run.font.size = Pt(12)
