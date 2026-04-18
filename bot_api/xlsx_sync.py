"""Конвертация SectionTable.content (custom_sheet) ↔ .xlsx (openpyxl)."""

from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, Tuple

from bot_api.sheet_utils import get_workbook_sheets

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None  # type: ignore


def excel_bytes_to_content(xlsx_bytes: bytes) -> Dict[str, Any]:
    """Читает .xlsx и возвращает content с ключом custom_sheet (v2)."""
    if not openpyxl:
        raise RuntimeError("openpyxl не установлен")
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), data_only=True)
    sheets_out = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            clean_row = ["" if cell is None else str(cell) for cell in row]
            rows.append(clean_row)
        sheets_out.append({"name": sheet_name, "data": rows})
    return {
        "custom_sheet": {
            "version": 2,
            "activeSheetIndex": 0,
            "sheets": sheets_out,
        }
    }


def content_to_excel_bytes(content: Dict[str, Any]) -> bytes:
    """Сериализует custom_sheet в .xlsx."""
    if not openpyxl:
        raise RuntimeError("openpyxl не установлен")
    sheets, _ = get_workbook_sheets(content if isinstance(content, dict) else {})
    wb = openpyxl.Workbook()
    default_ws = wb.active
    first = True

    if not sheets:
        default_ws.title = "Лист1"
        bio = BytesIO()
        wb.save(bio)
        return bio.getvalue()

    for sh in sheets:
        name = str(sh.get("name") or "Лист")[:31]
        data = sh.get("data") if isinstance(sh.get("data"), list) else []
        if first:
            ws = default_ws
            ws.title = name or "Лист1"
            first = False
        else:
            ws = wb.create_sheet(title=name or "Лист")
        for r_idx, row in enumerate(data, start=1):
            if not isinstance(row, list):
                row = list(row) if row else []
            for c_idx, val in enumerate(row, start=1):
                cell_val = None if val == "" or val is None else val
                ws.cell(row=r_idx, column=c_idx, value=cell_val)

    if first:
        default_ws.title = "Лист1"

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def build_custom_sheet_v2(sheets_payload: list) -> Dict[str, Any]:
    """Обёртка для листов в формате [{'name': ..., 'data': ...}, ...]."""
    return {
        "custom_sheet": {
            "version": 2,
            "activeSheetIndex": 0,
            "sheets": sheets_payload,
        }
    }


def content_sheet_count(content: Dict[str, Any]) -> Tuple[int, int]:
    """(число листов, число строк данных на первом листе)."""
    sheets, _ = get_workbook_sheets(content or {})
    if not sheets:
        return 0, 0
    data = sheets[0].get("data") if isinstance(sheets[0].get("data"), list) else []
    return len(sheets), len(data)
